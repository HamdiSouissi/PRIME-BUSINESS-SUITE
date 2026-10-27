"""Rendu de documents : substitution de champs, génération PDF/DOCX.

⚠️ Une image de signature collée dans un PDF n'est PAS une signature
électronique qualifiée : l'intégration d'un prestataire de signature reste à
brancher (champs prévus sur GeneratedDocument).
"""

import io
import re

from django.core.files.base import ContentFile
from django.utils import timezone

from apps.core.exceptions import BusinessError
from apps.core.services_pdf import html_vers_pdf

from .models import DocumentTemplate, GeneratedDocument, GeneratedDocumentStatus, TemplateStatus

VARIABLE_PATTERN = re.compile(r"\{\{\s*([\w.]+)\s*\}\}")


def variables_requises(template: DocumentTemplate) -> list[str]:
    return sorted(set(VARIABLE_PATTERN.findall(template.content or "")))


def rendre(template: DocumentTemplate, contexte: dict) -> str:
    """Substitue les variables. Toute variable manquante devient explicite."""
    texte = template.content or ""

    def remplace(match):
        cle = match.group(1)
        valeur = contexte.get(cle)
        if valeur is None or str(valeur) == "":
            return f"[{cle} — À COMPLÉTER]"
        return str(valeur)

    return VARIABLE_PATTERN.sub(remplace, texte)


def generer_document(template: DocumentTemplate, contexte: dict, user, label: str,
                     target_employee=None, target_client=None, ai_draft=False) -> GeneratedDocument:
    if template.status != TemplateStatus.VALIDE:
        raise BusinessError(
            "Seul un modèle validé juridiquement peut produire un document.",
            code="modele_non_valide",
        )
    contenu = rendre(template, contexte)
    document = GeneratedDocument.objects.create(
        company=template.company,
        template=template,
        label=label or f"{template.name} — {timezone.localdate():%Y-%m-%d}",
        content_snapshot=contenu,
        context_data=contexte,
        target_employee=target_employee,
        target_client=target_client,
        ai_draft=ai_draft,
        created_by=user,
    )
    return document


def generer_pdf(document: GeneratedDocument) -> GeneratedDocument:
    import html as html_module

    lignes = "".join(
        f"<p>{html_module.escape(ligne)}</p>" if ligne.strip() else "<br>"
        for ligne in document.content_snapshot.splitlines()
    )
    company = document.company
    page = f"""<!DOCTYPE html><html lang="fr"><head><meta charset="utf-8">
    <style>body{{font-family:Helvetica,Arial,sans-serif;font-size:11px;line-height:1.5;color:#111}}
    header{{border-bottom:2px solid #1e3a8a;padding-bottom:8px;margin-bottom:16px}}</style></head>
    <body><header><strong>{company.name}</strong><br>
    {company.address} {company.postal_code} {company.city}<br>SIRET : {company.siret}</header>
    {lignes}</body></html>"""
    pdf_bytes = html_vers_pdf(page)
    document.pdf.save(f"{document.label[:60]}.pdf", ContentFile(pdf_bytes), save=True)
    return document


def generer_docx(document: GeneratedDocument) -> bytes:
    from docx import Document as DocxDocument
    from docx.shared import Pt

    docx = DocxDocument()
    docx.add_heading(document.label, 0)
    para = docx.add_paragraph()
    run = para.add_run(document.company.name)
    run.bold = True
    docx.add_paragraph(document.company.address + " " + document.company.postal_code + " " + document.company.city)
    docx.add_paragraph("")
    for ligne in document.content_snapshot.splitlines():
        p = docx.add_paragraph(ligne)
        for run in p.runs:
            run.font.size = Pt(10)
    buffer = io.BytesIO()
    docx.save(buffer)
    return buffer.getvalue()


def valider_document(document: GeneratedDocument, user) -> GeneratedDocument:
    """Validation interne — relecture humaine obligatoire, a fortiori pour
    les brouillons générés par IA (information tracée)."""
    if document.status not in {GeneratedDocumentStatus.BROUILLON}:
        raise BusinessError("Document déjà validé.", code="transition_interdite")
    document.status = GeneratedDocumentStatus.VALIDE
    document.validated_by = user
    document.validated_at = timezone.now()
    document.save()
    return document
